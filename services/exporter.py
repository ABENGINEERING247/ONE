import io
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib import colors

def docx_bytes(title: str, markdown_content: str, style_standard: str = "Standard") -> bytes:
    """
    Generates DOCX bytes formatted according to the specified standard (Standard, IEEE, APA).
    """
    doc = Document()
    
    # Configure document margins and typography based on style
    section = doc.sections[0]
    
    if style_standard == "IEEE":
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.625)
        section.right_margin = Inches(0.625)
        font_name = "Times New Roman"
        font_size = 10
    elif style_standard == "APA":
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        font_name = "Times New Roman"
        font_size = 12
    else:  # Standard / General
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        font_name = "Calibri"
        font_size = 11

    # Apply normal style rules
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = font_name
    font.size = Pt(font_size)
    font.color.rgb = RGBColor(0, 0, 0)
    
    if style_standard == "APA":
        style_normal.paragraph_format.line_spacing = 2.0
        style_normal.paragraph_format.space_after = Pt(0)
    else:
        style_normal.paragraph_format.line_spacing = 1.15
        style_normal.paragraph_format.space_after = Pt(6)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(title)
    run_title.font.name = font_name
    run_title.font.size = Pt(24 if style_standard != "IEEE" else 18)
    run_title.bold = True
    
    # Process Markdown Content
    lines = markdown_content.splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h_run = h.add_run(stripped[2:])
            h_run.font.name = font_name
            h_run.bold = True
        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h_run = h.add_run(stripped[3:])
            h_run.font.name = font_name
            h_run.bold = True
        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h_run = h.add_run(stripped[4:])
            h_run.font.name = font_name
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        else:
            doc.add_paragraph(stripped)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def pdf_bytes(title: str, markdown_content: str, style_standard: str = "Standard") -> bytes:
    """
    Generates PDF bytes using ReportLab with custom styling according to standard.
    """
    buffer = io.BytesIO()
    
    # Setup page margins according to standard
    if style_standard == "IEEE":
        doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=45, rightMargin=45, topMargin=54, bottomMargin=54)
        body_font = "Times-Roman"
        body_size = 10
        leading = 12
    elif style_standard == "APA":
        doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=72, rightMargin=72, topMargin=72, bottomMargin=72)
        body_font = "Times-Roman"
        body_size = 12
        leading = 20
    else:
        doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
        body_font = "Helvetica"
        body_size = 10.5
        leading = 14

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName=body_font + "-Bold" if "Times" in body_font else "Helvetica-Bold",
        fontSize=20 if style_standard != "IEEE" else 16,
        leading=24,
        alignment=1, # Centered
        spaceAfter=15
    )
    
    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Heading2'],
        fontName=body_font + "-Bold" if "Times" in body_font else "Helvetica-Bold",
        fontSize=14,
        leading=18,
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName=body_font,
        fontSize=body_size,
        leading=leading,
        spaceAfter=8
    )

    story = [Paragraph(title, title_style), HRFlowable(width="100%", thickness=1, color=colors.grey, spaceAfter=15)]
    
    for line in markdown_content.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            clean_text = stripped.lstrip("#").strip()
            story.append(Paragraph(clean_text, h1_style))
        else:
            story.append(Paragraph(stripped, body_style))
            
    doc.build(story)
    return buffer.getvalue()


def latex_source(title: str, markdown_content: str, style_standard: str = "Standard") -> str:
    """
    Generates LaTeX document string.
    """
    doc_class = "\\documentclass[10pt, journal, compsoc]{IEEEtran}" if style_standard == "IEEE" else "\\documentclass[12pt]{article}"
    
    tex = f"""{doc_class}
\\usepackage[utf8]{{utf8}}
\\usepackage{{amsmath}}
\\usepackage{{hyperref}}

\\title{{{title}}}
\\author{{Generated by ONE AI Workspace}}
\\date{{\\today}}

\\begin{{document}}
\\maketitle

{markdown_content}

\\end{{document}}
"""
    return tex


def project_zip(title: str, markdown_content: str, sources: list, style_standard: str = "Standard") -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("manuscript.docx", docx_bytes(title, markdown_content, style_standard))
        zf.writestr("manuscript.pdf", pdf_bytes(title, markdown_content, style_standard))
        zf.writestr("manuscript.tex", latex_source(title, markdown_content, style_standard))
    return buffer.getvalue()
